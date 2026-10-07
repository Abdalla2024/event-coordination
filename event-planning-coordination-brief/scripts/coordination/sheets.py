"""Header-keyed parsing for the disclosed spreadsheet sources.

Fields are found by header name, never by column position, so reordered headers, reordered rows and
unrelated extra columns are accepted. Anything that cannot be interpreted safely is *held*: it is
reported as an issue and its row is left out of the interpreted records, so later stages see it as
missing rather than as a guessed value.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .util import safe_id

EXCEL_EPOCH = date(1899, 12, 30)


@dataclass
class Issue:
    id: str
    kind: str
    summary: str
    locator: str
    affected_ids: list[str] = field(default_factory=list)

    def as_record(self, source_id: str) -> dict:
        return {"id": self.id, "summary": self.summary, "evidence_ids": [source_id], "owner": None,
                "rationale": None, "kind": self.kind, "locator": self.locator,
                "affected_ids": self.affected_ids}


@dataclass
class TableParse:
    source_id: str
    sheet: str | None
    headers: list[dict]              # header text as captured, with its column letter
    records: dict[str, dict]         # interpreted rows keyed by id field
    observations: list[dict]         # one per interpreted row, with a cell-range locator
    issues: list[Issue]
    native_versions: list[str]

    @property
    def held(self) -> bool:
        return any(i.kind.startswith("table-") for i in self.issues)


class _Invalid(Exception):
    pass


def _norm(h) -> str:
    return str(h).strip().lower() if h is not None else ""


def _convert(value, spec: dict):
    t = spec["type"]
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if t == "string":
        return str(value).strip()
    if t == "enum":
        v = str(value).strip().lower()
        if v not in spec["values"]:
            raise _Invalid(f"value {value!r} is not one of {spec['values']}")
        return v
    if t == "int":
        if isinstance(value, bool):
            raise _Invalid(f"boolean {value!r} is not an integer")
        if isinstance(value, (int, float)):
            if float(value) != int(value):
                raise _Invalid(f"{value!r} is not a whole number")
            out = int(value)
        else:
            s = str(value).strip()
            if not s.isdigit():
                raise _Invalid(f"{value!r} is not a whole number")
            out = int(s)
        if "min" in spec and out < spec["min"]:
            raise _Invalid(f"{out} is below the minimum {spec['min']}")
        return out
    if t == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return (EXCEL_EPOCH + timedelta(days=int(value))).isoformat()
        try:
            return date.fromisoformat(str(value).strip()).isoformat()
        except ValueError:
            raise _Invalid(f"{value!r} is not an ISO date")
    if t == "datetime_tz":
        if isinstance(value, datetime):
            dt = value
        else:
            try:
                dt = datetime.fromisoformat(str(value).strip())
            except ValueError:
                raise _Invalid(f"{value!r} is not an ISO date-time")
        if dt.tzinfo is None:
            raise _Invalid(f"{value!r} has no explicit UTC offset")
        return dt.isoformat()
    raise ValueError(f"unknown field type {t}")


def parse_table(source_id: str, body: bytes, table: dict) -> TableParse:
    fields: dict = table["fields"]
    id_field: str = table["id_field"]
    issues: list[Issue] = []
    wb = load_workbook(io.BytesIO(body), data_only=True)
    visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]

    ws = wb[table["sheet"]] if table["sheet"] in wb.sheetnames else None
    if ws is None:
        if len(visible) == 1:
            ws = visible[0]
            issues.append(Issue(f"ISS-{source_id}-sheet-name", "sheet-renamed",
                                f"Expected tab '{table['sheet']}' not found; the only visible tab "
                                f"'{ws.title}' was read. Field checks still apply.", ws.title))
        else:
            issues.append(Issue(f"ISS-{source_id}-sheet", "table-missing-sheet",
                                f"Expected tab '{table['sheet']}' not found among {wb.sheetnames}; "
                                "no tab was guessed.", "workbook"))
            return TableParse(source_id, None, [], {}, [], issues, [])
    for other in wb.worksheets:
        if other is not ws and other.sheet_state != "visible":
            issues.append(Issue(f"ISS-{source_id}-hidden-tab-{safe_id(other.title)}", "hidden-tab",
                                f"Hidden tab '{other.title}' present; not read.", other.title))

    rows = list(ws.iter_rows(values_only=False))
    if not rows:
        issues.append(Issue(f"ISS-{source_id}-empty", "table-empty", "Sheet has no header row.", ws.title))
        return TableParse(source_id, ws.title, [], {}, [], issues, [])

    header_cells = rows[0]
    headers = [{"column": get_column_letter(c.column), "text": c.value} for c in header_cells
               if c.value is not None and str(c.value).strip()]
    by_name: dict[str, list[int]] = {}
    for c in header_cells:
        n = _norm(c.value)
        if n:
            by_name.setdefault(n, []).append(c.column)

    col_of: dict[str, int] = {}
    for name, spec in fields.items():
        cols = by_name.get(name, [])
        if len(cols) > 1:
            issues.append(Issue(f"ISS-{source_id}-dup-header-{name}", "table-duplicate-header",
                                f"Header '{name}' appears {len(cols)} times; meaning is ambiguous.",
                                f"'{ws.title}'!1:1"))
        elif not cols and spec["required"]:
            issues.append(Issue(f"ISS-{source_id}-missing-header-{name}", "table-missing-header",
                                f"Required header '{name}' not found. Renamed headers are not guessed.",
                                f"'{ws.title}'!1:1"))
        elif cols:
            col_of[name] = cols[0]
            letter = get_column_letter(cols[0])
            if ws.column_dimensions[letter].hidden:
                issues.append(Issue(f"ISS-{source_id}-hidden-col-{name}", "table-hidden-column",
                                    f"Column for '{name}' is hidden; visible meaning is unclear.",
                                    f"'{ws.title}'!{letter}:{letter}"))
    if any(i.kind.startswith("table-") for i in issues):
        return TableParse(source_id, ws.title, headers, {}, [], issues, [])

    last_col = get_column_letter(max(c.column for c in header_cells))
    candidates: list[tuple[int, str, dict]] = []
    invalid_ids: set[str] = set()
    for row in rows[1:]:
        r = row[0].row
        cells = {c.column: c.value for c in row}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in cells.values()):
            continue
        locator = f"'{ws.title}'!A{r}:{last_col}{r}"
        if ws.row_dimensions[r].hidden:
            issues.append(Issue(f"ISS-{source_id}-hidden-row-{r}", "hidden-row",
                                f"Row {r} is hidden; it is not counted as visible data.", locator))
            continue
        values, problems = {}, []
        for name, spec in fields.items():
            if name not in col_of:
                values[name] = None
                continue
            try:
                values[name] = _convert(cells.get(col_of[name]), spec)
            except _Invalid as exc:
                problems.append(f"{name}: {exc}")
                continue
            if values[name] is None and spec["required"]:
                problems.append(f"{name}: required value is blank")
        rid = values.get(id_field)
        if not problems and "end_not_before_start" in table.get("row_checks", []):
            if datetime.fromisoformat(values["end_at"]) < datetime.fromisoformat(values["start_at"]):
                problems.append("end_at is before start_at")
        if problems:
            tag = f"{safe_id(rid)}-r{r}" if rid else f"r{r}"
            issues.append(Issue(f"ISS-{source_id}-invalid-{tag}", "invalid-row",
                                f"Row {r} held: " + "; ".join(problems), locator,
                                [rid] if rid else []))
            if rid:
                invalid_ids.add(rid)
            continue
        candidates.append((r, locator, values))

    counts: dict[str, int] = {}
    for _, _, v in candidates:
        counts[v[id_field]] = counts.get(v[id_field], 0) + 1
    for rid in invalid_ids:
        if rid in counts:  # same identity also appears in an invalid row: ambiguous, hold both
            counts[rid] += 1
    records, observations = {}, []
    for r, locator, values in candidates:
        rid = values[id_field]
        if counts[rid] > 1:
            continue
        records[rid] = values
        observations.append({
            "id": f"OBS-{source_id}-{safe_id(rid)}",
            "summary": "; ".join(f"{k}={v}" for k, v in values.items() if v is not None),
            "locator": {"kind": "cell-range", "value": locator},
            "record_key": rid,
            "values": values,
        })
    for rid, n in counts.items():
        if n > 1:
            locs = [loc for _, loc, v in candidates if v[id_field] == rid]
            issues.append(Issue(f"ISS-{source_id}-duplicate-{safe_id(rid)}", "duplicate-id",
                                f"Identifier '{rid}' appears in {n} rows (including any invalid "
                                "row); none is chosen.",
                                " ; ".join(locs), [rid]))

    vfield = table.get("native_version_field")
    versions = sorted({v[vfield] for v in records.values() if vfield and v.get(vfield)})
    if len(versions) > 1:
        issues.append(Issue(f"ISS-{source_id}-mixed-versions", "mixed-record-versions",
                            f"Rows carry different {vfield} values: {versions}.", f"'{ws.title}'"))
    return TableParse(source_id, ws.title, headers, records, observations, issues, versions)
